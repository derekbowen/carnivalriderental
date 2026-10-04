#!/usr/bin/env python3
"""
Extract the company-account import contract from Carnival_Host_Import.xlsx into JSON.

  python3 scripts/company-import-extract.py <path/to/Carnival_Host_Import.xlsx>

Writes imports/company-accounts/source/workbook.json (gitignored: it holds private contacts).
Cells are copied verbatim; JSON cells stay strings here and are parsed and validated by
src/lib/imports/company-accounts.ts. The workbook's SHA-256 is recorded so every run
can say exactly which file it imported.
"""
import hashlib, json, os, sys

import openpyxl

USERS_HEADER = ["externalId", "importTier", "email (placeholder)", "firstName", "lastName", "displayName", "bio",
                "publicData (JSON)", "protectedData (JSON)", "privateData (JSON)", "metadata (JSON)"]
USERS_KEYS = ["externalId", "importTier", "emailPlaceholder", "firstName", "lastName", "displayName", "bio",
              "publicData", "protectedData", "privateData", "metadata"]


def main(path: str) -> None:
    digest = hashlib.sha256(open(path, "rb").read()).hexdigest()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    rows = list(wb["Sharetribe Users"].iter_rows(values_only=True))
    if list(rows[0][: len(USERS_HEADER)]) != USERS_HEADER:
        sys.exit(f"Sharetribe Users header changed: {rows[0]}")
    users = [dict(zip(USERS_KEYS, r[: len(USERS_KEYS)])) for r in rows[1:] if any(c is not None for c in r)]

    crows = list(wb["Companies"].iter_rows(values_only=True))
    h = list(crows[0])
    for col in ("Company ID", "Import tier"):
        if col not in h:
            sys.exit(f"Companies sheet is missing the {col!r} column")
    ci, ti = h.index("Company ID"), h.index("Import tier")
    companies = [{"companyId": r[ci], "importTier": r[ti]} for r in crows[1:] if r[ci] is not None]

    out = os.path.join("imports", "company-accounts", "source", "workbook.json")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "w") as f:
        json.dump({"source": {"file": os.path.basename(path), "sha256": digest}, "users": users, "companies": companies}, f)
    print(f"Wrote {out}: {len(users)} user rows, {len(companies)} companies (sha256 {digest[:12]}…)")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
