"""Generate research/operator-research-template.xlsx from contract/listing-contract.json.

Run: python3 scripts/make-research-template.py   (needs openpyxl)
The workbook is for collecting PUBLIC information about carnival ride operators and their rides.
It feeds our internal supplier list (call sheet), not public listings and not Sharetribe accounts.
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
contract = json.loads((ROOT / "contract/listing-contract.json").read_text())
categories = [c["id"] for c in contract["categories"]["items"]] + ["not-sure"]
states = [s.upper() for s in contract["optionSets"]["usStates"]["values"]]
statuses = ["researched", "contacted", "interested", "declined", "do-not-contact"]
yes_no = ["yes", "no", "unknown"]

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="7A1F2B")
REQ_FILL = PatternFill("solid", fgColor="B5442E")

COMPANIES = [
    # (header, width, required, note)
    ("company_id", 12, True, "CO-001, CO-002 … one per company, never reused"),
    ("company_name", 30, True, "As shown on their website"),
    ("website", 32, False, ""),
    ("phone", 16, False, "Public business number only"),
    ("email", 28, False, "Public business email only"),
    ("contact_name", 20, False, "Only if publicly listed or they told us"),
    ("city", 18, False, "Home base"),
    ("state", 8, True, "Home base state (dropdown)"),
    ("states_served", 24, False, "Comma list, e.g. TX, OK, LA — only what THEY say"),
    ("has_rental_website", 10, False, "yes / no / unknown"),
    ("status", 14, True, "researched until someone actually talks to them"),
    ("source_url", 40, True, "Where you found this company"),
    ("found_by", 14, True, "Your name"),
    ("date_found", 12, True, "YYYY-MM-DD"),
    ("notes", 40, False, ""),
]

RIDES = [
    ("ride_id", 12, True, "RD-001, RD-002 … one row per ride TYPE per company"),
    ("company_id", 12, True, "Must match a row on Companies"),
    ("ride_name_as_listed", 30, True, "Exactly what the company calls it"),
    ("our_category", 16, True, "Dropdown; not-sure is fine"),
    ("manufacturer", 18, False, "ONLY if the company states it — leave blank otherwise"),
    ("model", 18, False, "ONLY if stated — leave blank otherwise"),
    ("quantity", 10, False, "Only if stated"),
    ("advertised_price_usd", 14, False, "Only a price they publish; numbers only"),
    ("price_terms", 30, False, "What the price covers, e.g. 'per day, 4 hrs, 100-mile radius'"),
    ("source_url", 40, True, "Page where this ride is listed"),
    ("photo_page_url", 32, False, "Link only — do NOT download or reuse their photos"),
    ("notes", 40, False, ""),
]

INSTRUCTIONS = [
    "OPERATOR RESEARCH — how to fill this in",
    "",
    "Purpose: build our internal list of carnival ride companies and the rides they offer, so we know who to call when a customer requests a ride.",
    "This is an internal list. Nothing here is published on the website and nobody gets an account created for them.",
    "",
    "Rules",
    "1. Public information only (their website, Google Business profile, Facebook page, directories). Always fill source_url.",
    "2. Write what THEY say. If something isn't stated (manufacturer, size, price, insurance), leave it blank. Never guess.",
    "3. Do not download or copy their photos — paste the page link in photo_page_url instead.",
    "4. Status stays 'researched' until someone actually speaks with them. Only Derek decides who becomes a supplier.",
    "5. One row per company on Companies; one row per ride TYPE that company offers on Rides (two Ferris wheels = one row, quantity 2).",
    "6. Before calling or emailing any company, check with Derek — outreach is coordinated so companies hear one consistent message.",
    "7. Use do-not-contact if a company asks not to be contacted.",
    "",
    "Required columns have a red header. Dropdowns: state, status, our_category, has_rental_website.",
    "",
    "Category guide (pick the one a customer would name first)",
    "ferris-wheels — any Ferris / observation wheel",
    "carousels — carousels, merry-go-rounds",
    "swing-rides — wave swingers, chair swings, yo-yo",
    "kiddie-rides — rides meant mainly for small children (kiddie trains, teacups, mini coasters)",
    "family-rides — rides adults and kids ride together (scramblers, tilt-a-whirl, dragon coasters)",
    "thrill-rides — high-intensity rides (zipper, gravitron, drop towers, pendulums)",
    "not-sure — use it; we'll sort it later",
    "",
    "Good places to look: state fair/festival vendor lists, 'carnival ride rental [state]' searches, Outdoor Amusement Business Association member lists, company Facebook pages.",
]


def sheet(ws, cols, rows=500):
    for i, (name, width, required, note) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.font = HEAD
        cell.fill = REQ_FILL if required else HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        hint = ws.cell(row=2, column=i, value=note)
        hint.font = Font(italic=True, color="666666", size=9)
        hint.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[2].height = 42
    ws.freeze_panes = "A3"
    return {name: ws.cell(row=1, column=i).column_letter for i, (name, *_rest) in enumerate(cols, start=1)}


def dropdown(ws, col, ref, rows=1000):
    dv = DataValidation(type="list", formula1=ref, allow_blank=True, showErrorMessage=True)
    dv.add(f"{col}3:{col}{rows}")
    ws.add_data_validation(dv)


wb = Workbook()
guide = wb.active
guide.title = "How to fill this in"
guide.column_dimensions["A"].width = 120
for r, line in enumerate(INSTRUCTIONS, start=1):
    c = guide.cell(row=r, column=1, value=line)
    c.alignment = Alignment(wrap_text=True)
    if r == 1:
        c.font = Font(bold=True, size=14)
    elif line in ("Rules", "Category guide (pick the one a customer would name first)"):
        c.font = Font(bold=True)

lists = wb.create_sheet("Lists")
for col, (title, values) in enumerate(
    [("state", states), ("status", statuses), ("our_category", categories), ("yes_no", yes_no)], start=1
):
    lists.cell(row=1, column=col, value=title).font = Font(bold=True)
    for r, v in enumerate(values, start=2):
        lists.cell(row=r, column=col, value=v)


def list_ref(col_letter, n):
    return f"=Lists!${col_letter}$2:${col_letter}${n + 1}"


companies = wb.create_sheet("Companies", 1)
cc = sheet(companies, COMPANIES)
dropdown(companies, cc["state"], list_ref("A", len(states)))
dropdown(companies, cc["status"], list_ref("B", len(statuses)))
dropdown(companies, cc["has_rental_website"], list_ref("D", len(yes_no)))

rides = wb.create_sheet("Rides", 2)
rc = sheet(rides, RIDES)
dropdown(rides, rc["our_category"], list_ref("C", len(categories)))

out = ROOT / "research/operator-research-template.xlsx"
out.parent.mkdir(exist_ok=True)
wb.save(out)
print(f"wrote {out.relative_to(ROOT)}")
