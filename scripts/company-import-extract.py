#!/usr/bin/env python3
"""
Extract the company-account import contract from Carnival_Host_Import.xlsx into JSON.

  python3 scripts/company-import-extract.py <path/to/Carnival_Host_Import.xlsx> [provenance note]

The optional note records where the file came from (e.g. "Google Sheets XLSX export of Drive file <id>"),
because an export's SHA-256 is not the original upload's. Google exports drop trailing empty cells,
so every row is padded to its header width before mapping (a short row must read as nulls, not crash
or lose keys).

Writes imports/company-accounts/source/workbook.json (gitignored: it holds private contacts).
Cells are copied verbatim; JSON cells stay strings here and are parsed and validated by
src/lib/imports/company-accounts.ts. The workbook's SHA-256 is recorded so every run
can say exactly which file it imported.
"""
import hashlib, json, os, re, sys

import openpyxl

USERS_HEADER = ["externalId", "importTier", "email (placeholder)", "firstName", "lastName", "displayName", "bio",
                "publicData (JSON)", "protectedData (JSON)", "privateData (JSON)", "metadata (JSON)"]
USERS_KEYS = ["externalId", "importTier", "emailPlaceholder", "firstName", "lastName", "displayName", "bio",
              "publicData", "protectedData", "privateData", "metadata"]


_OOXML_ESC = re.compile(r"_x([0-9A-Fa-f]{4})_")


def cell(v):
    """OOXML escapes characters it cannot store as _xHHHH_ (Google exports do this for control
    characters); openpyxl's read-only mode leaves them escaped. Decode so text matches the source."""
    return _OOXML_ESC.sub(lambda m: chr(int(m.group(1), 16)), v) if isinstance(v, str) else v


def pad(r, n):
    r = tuple(cell(c) for c in r)
    return r + (None,) * (n - len(r)) if len(r) < n else r


def main(path: str, note: str | None = None) -> None:
    digest = hashlib.sha256(open(path, "rb").read()).hexdigest()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    rows = list(wb["Sharetribe Users"].iter_rows(values_only=True))
    if list(rows[0][: len(USERS_HEADER)]) != USERS_HEADER:
        sys.exit(f"Sharetribe Users header changed: {rows[0]}")
    users = [dict(zip(USERS_KEYS, pad(r, len(USERS_KEYS))[: len(USERS_KEYS)])) for r in rows[1:] if any(c is not None for c in r)]

    crows = list(wb["Companies"].iter_rows(values_only=True))
    h = list(crows[0])
    for col in ("Company ID", "Import tier"):
        if col not in h:
            sys.exit(f"Companies sheet is missing the {col!r} column")
    ci, ti = h.index("Company ID"), h.index("Import tier")
    crows = [pad(r, len(h)) for r in crows]
    companies = [{"companyId": r[ci], "importTier": r[ti]} for r in crows[1:] if r[ci] is not None]

    LISTINGS_HEADER = ["externalId", "authorExternalId", "title", "description", "lat", "lng", "image URL",
                       "publicData (JSON)", "privateData (JSON)", "metadata (JSON)"]
    LISTINGS_KEYS = ["externalId", "authorExternalId", "title", "description", "lat", "lng", "imageUrl",
                     "publicData", "privateData", "metadata"]
    lrows = list(wb["Sharetribe Listings"].iter_rows(values_only=True))
    if list(lrows[0][: len(LISTINGS_HEADER)]) != LISTINGS_HEADER:
        sys.exit(f"Sharetribe Listings header changed: {lrows[0]}")
    listings = [dict(zip(LISTINGS_KEYS, pad(r, len(LISTINGS_KEYS))[: len(LISTINGS_KEYS)])) for r in lrows[1:] if any(c is not None for c in r)]

    out = os.path.join("imports", "company-accounts", "source", "workbook.json")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "w") as f:
        json.dump({"source": {"file": os.path.basename(path), "sha256": digest, **({"provenance": note} if note else {})}, "users": users, "companies": companies, "listings": listings}, f)
    print(f"Wrote {out}: {len(users)} user rows, {len(companies)} companies, {len(listings)} listings (sha256 {digest[:12]}…)")


if __name__ == "__main__":
    if len(sys.argv) not in (2, 3):
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) == 3 else None)
